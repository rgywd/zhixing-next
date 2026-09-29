"""Structural validation inside the sandbox; visual inspection remains a separate step."""

import hashlib
import json
import sys
from pathlib import Path

path = Path(sys.argv[1])
assert path.is_file() and 0 < path.stat().st_size <= 20 * 1024 * 1024, (
    "Missing, empty or oversized artifact"
)
suffix = path.suffix.lower()
digest = hashlib.sha256(path.read_bytes()).hexdigest()
result = {
    "format": suffix,
    "sha256": digest,
    "checks": ["file_exists", "nonempty", "size_limit"],
    "visual_checked": False,
}
if suffix == ".xlsx":
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True)
    assert workbook.sheetnames, "Workbook has no worksheets"
    result["sheets"] = []
    for sheet in workbook:
        assert sheet.max_row <= 200000 and sheet.max_column <= 1000, (
            "Worksheet exceeds inspection budget"
        )
        errors = [
            cell.coordinate for row in sheet.iter_rows() for cell in row if cell.data_type == "e"
        ]
        assert not errors, "Spreadsheet error cells: " + ", ".join(errors[:20])
        result["sheets"].append(
            {"name": sheet.title, "rows": sheet.max_row, "columns": sheet.max_column}
        )
    result["checks"].extend(["workbook_reopened", "no_error_cells"])
    result["note"] = (
        "Formulas are not recalculated by openpyxl; use LibreOffice when cached results are required."
    )
    workbook.close()
elif suffix == ".docx":
    from docx import Document

    document = Document(path)
    result.update(paragraphs=len(document.paragraphs), tables=len(document.tables))
    result["checks"].append("document_reopened")
elif suffix == ".pptx":
    from pptx import Presentation

    presentation = Presentation(path)
    assert len(presentation.slides), "Presentation has no slides"
    result["slides"] = len(presentation.slides)
    result["checks"].append("presentation_reopened")
elif suffix == ".pdf":
    from pypdf import PdfReader

    document = PdfReader(path, strict=True)
    assert not document.is_encrypted and len(document.pages), "Unreadable PDF"
    result["pages"] = len(document.pages)
    result["checks"].append("pdf_reopened")
elif suffix in {".png", ".jpg", ".jpeg", ".webp", ".gif"}:
    from PIL import Image

    with Image.open(path) as image:
        image.verify()
    result["checks"].append("image_decoded")
elif suffix == ".json":
    json.loads(path.read_text())
    result["checks"].append("json_parsed")
elif suffix in {".txt", ".md", ".csv", ".html", ".xml", ".py", ".js", ".ts"}:
    path.read_text(encoding="utf-8")
    result["checks"].append("utf8_read")
else:
    result["note"] = "This format has only file existence and size checks."
assert hashlib.sha256(path.read_bytes()).hexdigest() == digest, (
    "Artifact changed during verification"
)
print(json.dumps(result, ensure_ascii=False))
